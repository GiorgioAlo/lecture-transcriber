"""FASE 4 — Generazione dei documenti Word (.docx).

Prende il testo finale di una lingua e produce un documento impaginato:
copertina + paragrafi, con le domande in corsivo e i timestamp in grassetto.
"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from . import util

TIMESTAMP_RE = re.compile(r"\[(\d{2}:\d{2}:\d{2})\]")
ALLCAPS_RE = re.compile(r"^[A-ZÀ-Ü0-9\s\-–—:]{3,}$")


def _sembra_titolo(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 70:
        return False
    # Una frase vera (finisce con . ? !) non e' un titolo, anche se urlata in
    # maiuscolo: evita che "OK." o "NON CAPISCO." diventino intestazioni.
    if s.endswith((".", "?", "!")):
        return False
    return s.endswith(":") or ALLCAPS_RE.match(s) is not None


def _aggiungi_copertina(doc, titolo, sottotitolo="", note=""):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titolo); r.bold = True; r.font.size = Pt(28)
    if sottotitolo:
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.add_run(sottotitolo).font.size = Pt(14)
    if note:
        p3 = doc.add_paragraph(); p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(note); r3.italic = True; r3.font.size = Pt(11)
    doc.add_page_break()


def _aggiungi_paragrafo(doc, testo):
    para = doc.add_paragraph()
    # Corsivo SOLO sulla frase che e' una domanda (finisce con "?"), non sul testo
    # che la segue: dopo il reflow un paragrafo contiene piu' frasi, quindi
    # isoliamo le domande frase per frase (_dividi_frasi gestisce le abbreviazioni).
    frasi = _dividi_frasi(testo) or [testo]
    for k, frase in enumerate(frasi):
        if k:
            para.add_run(" ")                                   # spazio tra le frasi
        is_q = frase.rstrip().rstrip(")]\"'”’»").endswith("?")  # domanda = finisce con "?"
        for i, part in enumerate(TIMESTAMP_RE.split(frase)):
            if not part:
                continue
            run = para.add_run(f"[{part}]" if i % 2 else part)
            if i % 2:
                run.bold = True
            if is_q:
                run.italic = True


# --- Reflow: ricuce i frammenti dell'ASR in paragrafi leggibili ------------
# La trascrizione (ASR) va a capo a OGNI pausa del parlato e il modello di testo
# trasforma quelle righe in paragrafi separati: cosi nel Word finiva un inciso
# per paragrafo (tipico nelle parti di dialogo/chiusura). Qui riuniamo i frammenti
# in prosa e li ri-raggruppiamo in paragrafi di lunghezza leggibile. Il TESTO non
# cambia (stesse parole): cambia solo come e' spezzato in righe.
REFLOW_TARGET_CHARS = 360          # lunghezza-bersaglio di un paragrafo

# Abbreviazioni dopo cui un punto NON chiude la frase: evita i falsi tagli tipo
# "Dr." o "e.g.". Confronto in minuscolo, con gli eventuali punti interni tolti.
_ABBR = {
    # inglese
    "dr", "mr", "mrs", "ms", "prof", "sr", "jr", "st", "vs", "etc", "eg", "ie",
    "cf", "al", "fig", "figs", "no", "nos", "inc", "ltd", "co", "corp", "eq",
    "eqs", "approx", "dept", "univ", "vol", "ch", "sec", "pp", "pg", "ca", "et",
    "viz", "esp", "us", "uk",
    # italiano
    "dott", "sig", "ing", "ecc", "es", "pag", "cfr", "art", "cap", "par", "n",
}


def _fine_frase(frase: str, prossimo: str) -> bool:
    """True se il punto in fondo a `frase` chiude DAVVERO una frase (non e'
    un'abbreviazione ne' un'iniziale puntata). `prossimo` = testo che segue."""
    m = re.search(r"([A-Za-zÀ-ü.'’‘]+)[.!?]+$", frase)
    ultima = re.sub(r"[.'’‘]", "", m.group(1) if m else "").lower()
    if ultima in _ABBR:                 # Dr., Prof., e.g., ecc.
        return False
    if len(ultima) == 1:                # iniziale puntata ("A.", o la "g" di "e.g.")
        return False
    p = prossimo.lstrip()
    if p and not (p[0].isupper() or p[0].isdigit() or p[0] in "\"'“‘([«"):
        return False                    # cio' che segue non inizia "da frase": punto interno
    return True


def _dividi_frasi(testo: str) -> list[str]:
    """Spezza il testo in frasi su . ! ? , ricucendo i tagli falsi (_fine_frase)."""
    parti = re.split(r"([.!?]+)(\s+)", testo)   # [corpo, punt, spazio, corpo, ...]
    frasi, cur = [], ""
    for i in range(0, len(parti), 3):
        cur += parti[i] + (parti[i + 1] if i + 1 < len(parti) else "")
        punt = parti[i + 1] if i + 1 < len(parti) else ""
        prossimo = parti[i + 3] if i + 3 < len(parti) else ""
        if punt and _fine_frase(cur, prossimo):
            frasi.append(cur.strip()); cur = ""
        elif i + 2 < len(parti):
            cur += parti[i + 2]                 # spazio: la frase continua
    if cur.strip():
        frasi.append(cur.strip())
    return frasi


def _reflow(righe: list[str]) -> list[str]:
    """Unisce righe-frammento in un flusso e lo raggruppa in paragrafi di ~TARGET
    caratteri (almeno 2 frasi l'uno), tagliando solo a fine frase."""
    flusso = re.sub(r"\s+", " ", " ".join(righe)).strip()
    if not flusso:
        return []
    paras, cur, n = [], "", 0
    for f in _dividi_frasi(flusso):
        cur = (cur + " " + f).strip() if cur else f
        n += 1
        if len(cur) >= REFLOW_TARGET_CHARS and n >= 2:
            paras.append(cur); cur, n = "", 0
    if cur:
        paras.append(cur)
    return paras


def genera_docx(testo: str, path: Path, titolo: str, note: str) -> Path:
    """Scrive il documento Word e ritorna il percorso creato."""
    doc = Document()
    _aggiungi_copertina(doc, titolo, f"Generato il {datetime.now():%d/%m/%Y %H:%M}", note)
    testo = (testo or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    # Accumuliamo la prosa e la mandiamo al reflow; un titolo o un elenco puntato
    # interrompe il flusso (viene emesso cosi com'e'), poi la prosa riprende.
    prosa: list[str] = []

    def _svuota_prosa():
        for p in _reflow(prosa):
            _aggiungi_paragrafo(doc, p)
        prosa.clear()

    for riga in (testo.split("\n") if testo else []):
        s = riga.strip()
        if not s:
            continue
        if _sembra_titolo(s):
            _svuota_prosa()
            doc.add_paragraph(s.rstrip(":"), style="Heading 2")
        elif s.startswith(("• ", "- ", "* ")):
            _svuota_prosa()
            doc.add_paragraph(s.lstrip("•-* ").strip(), style="List Bullet")
        else:
            prosa.append(s)
    _svuota_prosa()
    # Salvataggio atomico: scrive un .docx temporaneo e poi lo rinomina, cosi un
    # kill a meta' (Annulla immediato) non lascia un documento corrotto.
    tmp, finale = util.percorso_atomico(path)
    doc.save(tmp)
    util.concludi_atomico(tmp, finale)
    return path
